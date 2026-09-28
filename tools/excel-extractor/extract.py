"""Extrae fórmulas, valores calculados, validaciones y nombres definidos de un libro de Excel.

Uso:
    uv run --with openpyxl python tools/excel-extractor/extract.py <libro.xlsx> <carpeta_salida>

Genera en la carpeta de salida:
    <libro>.cells.json   una entrada por celda no vacía: hoja, celda, fórmula, valor calculado
    <libro>.meta.json    hojas, dimensiones, nombres definidos, validaciones de datos, formatos condicionales
    <libro>.formulas.md  inventario legible de fórmulas agrupadas por hoja

El libro se abre dos veces: una para leer las fórmulas y otra con data_only=True para leer
los valores que Excel guardó en caché la última vez que recalculó.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# Referencias a otras hojas: 'Hoja con espacios'!A1 o Hoja!A1
SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-zÀ-ÿ_][\wÀ-ÿ\.]*))!\$?[A-Z]{1,3}\$?\d+")


def _formula_text(value):
    """Devuelve el texto de la fórmula, incluidas las fórmulas de matriz."""
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


def _jsonable(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def extract(path: Path, out_dir: Path) -> None:
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)

    cells = []
    meta = {"file": path.name, "sheets": [], "defined_names": {}}

    for name, dn in wb_f.defined_names.items():
        meta["defined_names"][name] = dn.attr_text

    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        formulas = 0
        cached_missing = 0
        deps = set()
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                formula = _formula_text(c.value)
                cached = wv[c.coordinate].value
                if formula:
                    formulas += 1
                    if cached is None:
                        cached_missing += 1
                    for quoted, bare in SHEET_REF.findall(formula):
                        ref = quoted or bare
                        if ref != ws.title:
                            deps.add(ref)
                cells.append(
                    {
                        "sheet": ws.title,
                        "cell": c.coordinate,
                        "formula": formula,
                        "value": _jsonable(cached if formula else c.value),
                        "fill": (c.fill.fgColor.rgb if c.fill and c.fill.fgColor and c.fill.fgColor.type == "rgb" else None),
                    }
                )

        validations = []
        if ws.data_validations:
            for dv in ws.data_validations.dataValidation:
                validations.append({"ranges": str(dv.sqref), "type": dv.type, "formula1": dv.formula1})

        cond = []
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for r in rules:
                cond.append({"range": str(rng.sqref), "type": r.type, "formula": list(r.formula or [])})

        meta["sheets"].append(
            {
                "name": ws.title,
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_col": get_column_letter(ws.max_column),
                "formula_count": formulas,
                "formulas_without_cached_value": cached_missing,
                "depends_on": sorted(deps),
                "data_validations": validations,
                "conditional_formats": cond,
                "tab_color": ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else None,
            }
        )

    stem = path.stem
    (out_dir / f"{stem}.cells.json").write_text(json.dumps(cells, ensure_ascii=False, indent=1))
    (out_dir / f"{stem}.meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1))

    lines = [f"# Inventario de fórmulas: {path.name}", ""]
    for s in meta["sheets"]:
        lines.append(f"## {s['name']}")
        lines.append("")
        lines.append(
            f"Dimensiones {s['dimensions']}, {s['formula_count']} fórmulas, "
            f"{s['formulas_without_cached_value']} sin valor en caché. "
            f"Depende de: {', '.join(s['depends_on']) or 'ninguna hoja'}."
        )
        lines.append("")
        lines.append("| Celda | Fórmula | Valor en caché |")
        lines.append("|---|---|---|")
        for c in cells:
            if c["sheet"] == s["name"] and c["formula"]:
                f = c["formula"].replace("|", "\\|")
                lines.append(f"| {c['cell']} | `{f}` | {c['value']} |")
        lines.append("")
    (out_dir / f"{stem}.formulas.md").write_text("\n".join(lines))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    out = Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    extract(Path(sys.argv[1]), out)
