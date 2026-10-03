"""Extrae un caso de prueba de oro de un libro ya recalculado en Excel (ver recalc.py).

Uso:
    uv run --with openpyxl python tools/excel-extractor/golden.py <libro_recalculado.xlsx>
        <carpeta_del_caso> --case c3-plantilla-vacia --template Plantilla_Asesoria_Financiera.xlsx
        [--forbid terminos.txt]

Genera en la carpeta del caso:
    case.json      metadatos: caso, plantilla, fecha de corte, conteos
    inputs.json    celdas de entrada (crema o ámbar) con un valor escrito, más las entradas
                   adicionales del perfil de la plantilla, por hoja y celda
    expected.json  todas las celdas con fórmula y el valor que Excel calculó, por hoja y celda

Falla si:
    - la fecha de corte es una fórmula (por ejemplo =TODAY()) en lugar de una fecha fija;
    - alguna fórmula tiene un error de Excel (#DIV/0!, #N/A, ...) o no tiene valor calculado;
    - con --forbid, algún término del archivo (uno por línea, sin distinguir mayúsculas) aparece
      en cualquier celda del libro: textos, fórmulas, resultados o comentarios. El archivo de
      términos vive fuera del repositorio.

Convenciones de los valores: números como JSON (doble precisión, igual que Excel), fechas como
"AAAA-MM-DD", textos vacíos como "", verdadero y falso como true y false.
"""

import argparse
import json
import re
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree

import openpyxl

INPUT_FILLS = {"FFFFF8E6", "FFFDE9B8"}  # crema (editable) y ámbar (por confirmar)
EXCEL_ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!"}
# Perfil de cada plantilla: celda de la fecha de corte y celdas que no son crema pero cambian
# por cliente (la moneda base de Listas!M2 cambió a EUR en el caso de España, hallazgo H-15).
TEMPLATES = {
    "Plantilla_Asesoria_Financiera.xlsx": {"cutoff": "Supuestos!C12", "extra_inputs": ["Listas!M2"]},
    "Plantilla_Creditos.xlsx": {"cutoff": "Datos!C7", "extra_inputs": []},
}
NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def is_formula(value) -> bool:
    return (isinstance(value, str) and value.startswith("=")) or hasattr(value, "text")


def is_input(cell) -> bool:
    fg = cell.fill.fgColor if cell.fill else None
    return fg is not None and fg.type == "rgb" and fg.rgb in INPUT_FILLS


def normalize(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def empty_text_formulas(path: Path) -> dict[str, set[str]]:
    """Celdas con fórmula cuyo resultado es texto vacío (t="str" y <v/>), por hoja.

    openpyxl las devuelve como None, igual que una celda sin valor calculado; el XML las distingue.
    """
    result: dict[str, set[str]] = {}
    with zipfile.ZipFile(path) as z:
        workbook = ElementTree.fromstring(z.read("xl/workbook.xml"))
        rels = ElementTree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            target = targets[sheet.get(f"{{{NS['r']}}}id")].lstrip("/")
            part = target if target.startswith("xl/") else f"xl/{target}"
            cells = set()
            for c in ElementTree.fromstring(z.read(part)).iter(f"{{{NS['m']}}}c"):
                v = c.find("m:v", NS)
                if c.get("t") == "str" and c.find("m:f", NS) is not None and (v is None or not v.text):
                    cells.add(c.get("r"))
            result[sheet.get("name")] = cells
    return result


def load_forbidden(path: Path | None) -> list[str]:
    if path is None:
        return []
    terms = [line.strip() for line in path.read_text(encoding="utf-8").splitlines()]
    return [t.casefold() for t in terms if t and not t.startswith("#")]


def extract(path: Path, profile: dict, forbidden: list[str]):
    cutoff_cell = profile["cutoff"]
    extra = {tuple(ref.split("!")) for ref in profile["extra_inputs"]}
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    empty_text = empty_text_formulas(path)
    inputs: dict[str, dict] = {}
    expected: dict[str, dict] = {}
    problems: list[str] = []

    for ws in wb_f.worksheets:
        values = wb_v[ws.title]
        sheet_inputs: dict = {}
        sheet_expected: dict = {}
        for row in ws.iter_rows():
            for cell in row:
                ref = f"{ws.title}!{cell.coordinate}"
                raw = cell.value
                if raw is None and cell.comment is None:
                    continue
                texts = [str(getattr(raw, "text", raw))]
                if cell.comment is not None:
                    texts.append(cell.comment.text)
                folded = " ".join(texts).casefold()
                problems += [f"{ref}: contiene un término prohibido" for t in forbidden if t in folded]
                if raw is None:
                    continue
                if is_formula(raw):
                    value = values[cell.coordinate].value
                    if value is None:
                        if cell.coordinate in empty_text.get(ws.title, set()):
                            value = ""
                        else:
                            problems.append(f"{ref}: fórmula sin valor calculado")
                    elif isinstance(value, str) and value in EXCEL_ERRORS:
                        problems.append(f"{ref}: error {value}")
                    elif isinstance(value, str):
                        folded = value.casefold()
                        problems += [f"{ref}: resultado con un término prohibido" for t in forbidden if t in folded]
                    sheet_expected[cell.coordinate] = normalize(value)
                elif is_input(cell) or (ws.title, cell.coordinate) in extra:
                    sheet_inputs[cell.coordinate] = normalize(raw)
        if sheet_inputs:
            inputs[ws.title] = sheet_inputs
        if sheet_expected:
            expected[ws.title] = sheet_expected

    sheet, coord = cutoff_cell.split("!")
    cutoff = wb_f[sheet][coord].value
    if is_formula(cutoff) or not isinstance(cutoff, datetime):
        problems.append(f"{cutoff_cell}: la fecha de corte debe ser una fecha fija, no {cutoff!r}")
        cutoff_iso = None
    else:
        cutoff_iso = normalize(cutoff)
    return inputs, expected, cutoff_iso, problems


def write_json(path: Path, data) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--case", required=True, help="Identificador del caso, por ejemplo c2-espana")
    parser.add_argument("--template", required=True, choices=sorted(TEMPLATES), help="Plantilla de origen")
    parser.add_argument("--forbid", type=Path, help="Términos identificables, uno por línea")
    args = parser.parse_args()

    if not re.fullmatch(r"c\d+-[a-z0-9-]+", args.case):
        print("El caso debe tener la forma c<n>-<nombre>, por ejemplo c3-plantilla-vacia", file=sys.stderr)
        return 1

    profile = TEMPLATES[args.template]
    inputs, expected, cutoff, problems = extract(args.workbook, profile, load_forbidden(args.forbid))
    if problems:
        print(f"{len(problems)} problemas; no se escribió nada:", file=sys.stderr)
        for p in problems[:50]:
            print(f"  {p}", file=sys.stderr)
        return 1

    args.out_dir.mkdir(parents=True, exist_ok=True)
    count = lambda d: sum(len(cells) for cells in d.values())  # noqa: E731
    write_json(args.out_dir / "case.json", {
        "case": args.case,
        "template": args.template,
        "cutoffCell": profile["cutoff"],
        "cutoffDate": cutoff,
        "inputCells": count(inputs),
        "formulaCells": count(expected),
    })
    write_json(args.out_dir / "inputs.json", inputs)
    write_json(args.out_dir / "expected.json", expected)
    print(f"{args.case}: {count(inputs)} entradas y {count(expected)} fórmulas en {args.out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
